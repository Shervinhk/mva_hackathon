"""Create the public biology/methods guide from the public Markdown document.

This script never opens patient inputs, private analysis output or private logs.
Run with the bundled document runtime (python-docx); render and inspect afterward.
"""
from __future__ import annotations
import argparse
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE

ROOT = Path(__file__).resolve().parents[1]


def inline(paragraph, text):
    pattern = r'(\[[^\]]+\]\(https?://[^)]+\)|`[^`]+`|\*\*[^*]+\*\*)'
    for part in re.split(pattern, text):
        if not part:
            continue
        link = re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)', part)
        if link:
            relation = paragraph.part.relate_to(link[2], RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
            hyperlink = OxmlElement('w:hyperlink')
            hyperlink.set(qn('r:id'), relation)
            run = OxmlElement('w:r')
            properties = OxmlElement('w:rPr')
            color = OxmlElement('w:color')
            color.set(qn('w:val'), '1E5675')
            properties.append(color)
            run.append(properties)
            value = OxmlElement('w:t')
            value.text = link[1]
            run.append(value)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
        else:
            is_code = part.startswith('`') and part.endswith('`')
            is_bold = part.startswith('**') and part.endswith('**')
            run = paragraph.add_run(part[1:-1] if is_code else part[2:-2] if is_bold else part)
            if is_code:
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
            run.bold = is_bold


def table(document, lines):
    rows = [[part.strip() for part in line.strip().strip('|').split('|')] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r':?-+:?', value) for value in row)]
    result = document.add_table(rows=len(rows), cols=len(rows[0]))
    result.autofit = False
    widths = [1.35, 2.08, 3.47] if len(rows[0]) == 3 and rows[0][0] == 'Area' else [1.7, 2.4, 2.8]
    for column, width in zip(result.columns, widths):
        column.width = Inches(width)
    props = result._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = OxmlElement('w:' + edge)
        for name, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            element.set(qn('w:' + name), value)
        borders.append(element)
    props.append(borders)
    for i, values in enumerate(rows):
        row = result.rows[i]
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement('w:cantSplit'))
        if i == 0:
            row_properties.append(OxmlElement('w:tblHeader'))
        for j, text in enumerate(values):
            cell = row.cells[j]
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ('top', 'left', 'bottom', 'right'):
                margin = OxmlElement('w:' + edge)
                margin.set(qn('w:w'), '100')
                margin.set(qn('w:type'), 'dxa')
                margins.append(margin)
            properties.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '234D62' if i == 0 else 'F0F4F6' if i % 2 else 'FFFFFF')
            properties.append(shade)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.08
            inline(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(10.5)
                run.font.color.rgb = RGBColor.from_string('FFFFFF' if i == 0 else '000000')
                run.bold = i == 0 or run.bold
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=ROOT / 'outputs/documents/Understanding_MVA_Track1_v3.docx')
    args = parser.parse_args()
    source = (ROOT / 'docs/TRACK1_V3.md').read_text()
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.8)
    for name in ('Normal', 'Body Text', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Bullet', 'List Number'):
        style = document.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.widow_control = True
    normal = document.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.13
    normal.paragraph_format.space_after = Pt(7)
    for name in ('List Bullet', 'List Number'):
        document.styles[name].font.size = Pt(11)
        document.styles[name].paragraph_format.space_after = Pt(5)
    document.styles['Title'].font.size = Pt(28)
    document.styles['Title'].paragraph_format.space_after = Pt(12)
    document.styles['Heading 1'].font.size = Pt(15)
    document.styles['Heading 1'].paragraph_format.space_before = Pt(16)
    document.styles['Heading 1'].paragraph_format.space_after = Pt(7)
    document.styles['Heading 1'].paragraph_format.keep_with_next = True
    document.styles['Heading 2'].font.size = Pt(12)
    # Word's bundled base template can carry a blue title border. Remove
    # paragraph-border residue explicitly while leaving table borders intact.
    for border in list(document.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    document.add_paragraph('Understanding Track 1 v3', style='Title')
    document.add_paragraph('Biology data analysis and submission review', style='Subtitle')
    document.add_paragraph('The central task is to identify a genetic explanation that fits the child’s symptoms and inheritance. '
        'Version 3 makes the evidence easier to compare and review. A high rank is a starting point for that review; '
        'it does not establish that a variant causes the condition.')
    lines = source.splitlines()[1:]
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith('This is a research workflow for ranking'):
            i += 1
            continue
        if line.startswith('```'):
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.0
                run = paragraph.add_run(lines[i])
                run.font.name, run.font.size = 'Consolas', Pt(9)
                i += 1
            i += 1
            continue
        if line.startswith('|'):
            content = []
            while i < len(lines) and lines[i].startswith('|'):
                content.append(lines[i])
                i += 1
            table(document, content)
            continue
        heading = re.match(r'^(#{2,3}) (.+)$', line)
        if heading:
            text = re.sub(r'[^A-Za-z0-9 ]+', ' ', heading[2])
            document.add_paragraph(' '.join(text.split()), style='Heading 1' if len(heading[1]) == 2 else 'Heading 2')
        else:
            bullet = re.match(r'^- (.+)$', line)
            number = re.match(r'^\d+\. (.+)$', line)
            paragraph = document.add_paragraph(style='List Bullet' if bullet else 'List Number' if number else 'Normal')
            inline(paragraph, bullet[1] if bullet else number[1] if number else line)
        i += 1
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Track 1 v3  |  ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0, 0, 0)
    document.core_properties.title = 'Understanding Track 1 v3'
    document.core_properties.subject = 'Biology and local variant prioritization workflow'
    document.core_properties.author = 'MVA Track 1 research project'
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    print('Public biology and methods guide created.')


if __name__ == '__main__':
    main()
